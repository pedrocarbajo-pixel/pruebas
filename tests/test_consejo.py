import json
from pathlib import Path
import threading
import time
from types import SimpleNamespace as NS

import pytest

from council import privacidad
from council.consejo import Consejo
from council.informe import guardar
from council.providers import MiembroClaude, MiembroOpenAI, formatear_con_fuentes
from council.servicio import ejecutar
from council.web import crear_app


class MiembroFalso:
    def __init__(self, nombre):
        self.nombre = nombre
        self.modelo = f"{nombre.lower()}-falso"
        self.llamadas = []
        self._cerrojo = threading.Lock()

    def preguntar(self, sistema, mensaje):
        with self._cerrojo:
            self.llamadas.append(mensaje)
        if "INDICACIONES DEL USUARIO" in mensaje:
            return f"## Respuesta final\nVersión corregida por {self.nombre}"
        if "PRESIDENTE" in mensaje:
            return f"## Respuesta final\nSíntesis de {self.nombre}: {mensaje[-60:]}"
        if "Respuesta de Miembro" in mensaje:
            return f"Revisión de {self.nombre}"
        return f"Borrador de {self.nombre}"


def fabricas_falsas():
    return {clave: (lambda n=nombre, nivel="estandar": MiembroFalso(n))
            for clave, nombre in [("claude", "Claude"), ("chatgpt", "ChatGPT"),
                                  ("perplexity", "Perplexity")]}


def consejo_de_tres():
    miembros = [MiembroFalso(n) for n in ("Claude", "ChatGPT", "Perplexity")]
    return miembros, Consejo(miembros, registro=lambda *_: None)


def test_deliberacion_con_tres_miembros():
    (a, b, c), consejo = consejo_de_tres()
    r = consejo.deliberar("Revisa", "Contrato X", modo="contrato_revisar")
    assert set(r.borradores) == {"Claude", "ChatGPT", "Perplexity"}
    assert set(r.revisiones) == {"Claude", "ChatGPT", "Perplexity"}
    assert "Síntesis de Claude" in r.sintesis
    # Cada miembro: 1 borrador + 1 revisión; el presidente, además, la síntesis (3 + 3 + 1 = 7).
    assert (len(a.llamadas), len(b.llamadas), len(c.llamadas)) == (3, 2, 2)
    # Cada revisor recibe las respuestas de los otros dos, no la propia, con autores anonimizados.
    revision_de_a = next(m for m in a.llamadas if "Respuesta de Miembro" in m)
    assert "Borrador de ChatGPT" in revision_de_a and "Borrador de Perplexity" in revision_de_a
    assert "Borrador de Claude" not in revision_de_a
    assert "### Respuesta de Miembro B" in revision_de_a and "Miembro A" not in revision_de_a
    # El presidente no ve los nombres reales de los autores.
    sintesis = a.llamadas[-1]
    assert "### Borrador de Miembro C" in sintesis
    assert "### Borrador de Perplexity" not in sintesis


@pytest.mark.parametrize("modo", ["informe_tributario", "procedimiento_tributario", "dictamen_juridico"])
def test_modos_tributarios_y_juridicos(modo):
    (a, _, _), consejo = consejo_de_tres()
    consejo.deliberar("Tarea", modo=modo)
    assert "Modo:" in a.llamadas[0]


def test_presidente_configurable_y_validaciones():
    a, b = MiembroFalso("Claude"), MiembroFalso("ChatGPT")
    r = Consejo([a, b], presidente=b, registro=lambda *_: None).deliberar("T")
    assert r.presidente == "ChatGPT" and "Síntesis de ChatGPT" in r.sintesis
    with pytest.raises(ValueError):
        Consejo([a])
    with pytest.raises(ValueError):
        Consejo([a, b]).deliberar("T", modo="inexistente")


def test_fuentes_de_perplexity():
    respuesta = NS(
        output_text="Según el art. 66 LGT…",
        output=[
            NS(type="search_results", results=[NS(url="https://www.boe.es/a", title="LGT")]),
            NS(type="message", content=[NS(annotations=[
                NS(url="https://www.boe.es/a", title="LGT"),
                NS(url="https://petete.tributos.hacienda.gob.es/b", title="Consulta DGT"),
                NS(url=None, title=None),
            ])]),
        ],
    )
    texto = formatear_con_fuentes(respuesta)
    assert texto.startswith("Según el art. 66 LGT…")
    assert "1. [LGT](https://www.boe.es/a)" in texto
    assert "2. [Consulta DGT](https://petete.tributos.hacienda.gob.es/b)" in texto
    assert texto.count("boe.es/a") == 1


def test_seudonimizacion_reversible():
    texto = ("D. Juan, DNI 12345678Z, IBAN ES91 2100 0418 4502 0005 1332, "
             "correo juan@ejemplo.es, tel. 612 345 678. Repite DNI 12345678Z. "
             "Sociedad B12345674, NIE X1234567L.")
    seudo, mapa = privacidad.seudonimizar(texto)
    for dato in ["12345678Z", "ES91", "juan@ejemplo.es", "612 345 678", "B12345674", "X1234567L"]:
        assert dato not in seudo
    assert seudo.count("[DNI_1]") == 2
    assert privacidad.restaurar(seudo, mapa) == texto


def test_guardar_restaura_datos(tmp_path):
    a, b = MiembroFalso("Claude"), MiembroFalso("ChatGPT")
    r = Consejo([a, b], registro=lambda *_: None).deliberar("Cliente [EMAIL_1]")
    mapa = {"[EMAIL_1]": "x@y.es"}
    md, js = guardar(r, tmp_path, restaurar=lambda t: privacidad.restaurar(t, mapa))
    assert "x@y.es" in md.read_text(encoding="utf-8")
    assert json.loads(js.read_text(encoding="utf-8"))["tarea"] == "Cliente x@y.es"


def test_servicio_seudonimiza_antes_de_enviar(tmp_path):
    enviados = []

    class Espia(MiembroFalso):
        def preguntar(self, sistema, mensaje):
            enviados.append(mensaje)
            return super().preguntar(sistema, mensaje)

    fabricas = {"claude": lambda nivel: Espia("Claude"), "chatgpt": lambda nivel: Espia("ChatGPT")}
    salida = ejecutar("Cliente con DNI 12345678Z", "IBAN ES9121000418450200051332",
                      seudonimizar=True, carpeta=tmp_path, registro=lambda *_: None,
                      fabricas=fabricas)
    assert all("12345678Z" not in m and "ES9121" not in m for m in enviados)
    assert "12345678Z" in salida.informe.read_text(encoding="utf-8")
    assert salida.datos_seudonimizados == 2 and salida.mapa.exists()
    with pytest.raises(ValueError):
        ejecutar("T", miembros=["claude", "chatgpt"], presidente="perplexity",
                 fabricas=fabricas, carpeta=tmp_path)


def test_web_flujo_completo(tmp_path):
    cliente = crear_app(tmp_path, fabricas=fabricas_falsas()).test_client()
    conf = cliente.get("/api/configuracion").get_json()
    assert {m["id"] for m in conf["miembros"]} == {"claude", "chatgpt", "perplexity"}
    assert "informe_tributario" in {m["id"] for m in conf["modos"]}
    assert cliente.get("/").status_code == 200

    # Validaciones
    assert cliente.post("/api/consultas", data={"tarea": ""}).status_code == 400
    assert cliente.post("/api/consultas", data={
        "tarea": "T", "miembros": ["claude"], "presidente": "claude"}).status_code == 400

    r = cliente.post("/api/consultas", data={
        "tarea": "Analiza la liquidación", "modo": "procedimiento_tributario",
        "miembros": ["claude", "chatgpt", "perplexity"], "presidente": "perplexity",
        "seudonimizar": "on", "nivel": "extremo",
    })
    assert r.status_code == 202
    id_consulta = r.get_json()["id"]
    for _ in range(100):
        estado = cliente.get(f"/api/consultas/{id_consulta}").get_json()
        if estado["estado"] != "en_curso":
            break
        time.sleep(0.05)
    assert estado["estado"] == "terminada", estado
    assert "Síntesis de Perplexity" in estado["informe"]
    assert "**Nivel:** Extremo" in estado["informe"]
    assert any("Fase 3/3" in m for m in estado["mensajes"])
    descarga = cliente.get(f"/api/consultas/{id_consulta}/registro")
    assert descarga.status_code == 200 and json.loads(descarga.data)["modo"] == "procedimiento_tributario"
    assert cliente.get("/api/consultas/inexistente").status_code == 404


class ErrorPeticion(Exception):
    status_code = 400


class RespuestasOpenAIFalsas:
    """Rechaza el esfuerzo «max» como lo haría un modelo que no lo admite."""

    def __init__(self):
        self.peticiones = []

    def create(self, **parametros):
        self.peticiones.append(parametros)
        if parametros.get("reasoning", {}).get("effort") == "max":
            raise ErrorPeticion("effort no admitido")
        return NS(output_text="ok", output=[])


def test_extremo_openai_rebaja_esfuerzo_no_admitido(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "clave-falsa")
    monkeypatch.delenv("OPENAI_MODEL_EXTREMO", raising=False)
    miembro = MiembroOpenAI(nivel="extremo")
    falsas = RespuestasOpenAIFalsas()
    miembro._cliente = NS(responses=falsas)
    assert miembro.modelo == "gpt-6-astra"
    assert miembro.preguntar("s", "m") == "ok" and miembro.preguntar("s", "m") == "ok"
    esfuerzos = [p["reasoning"]["effort"] for p in falsas.peticiones]
    assert esfuerzos == ["max", "xhigh", "xhigh"]  # el nivel rebajado se mantiene
    assert miembro.descripcion() == "gpt-6-astra · esfuerzo xhigh · búsqueda web"


def test_estandar_openai_sin_esfuerzo(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "clave-falsa")
    monkeypatch.setenv("OPENAI_MODEL", "modelo-normal")
    monkeypatch.setenv("BUSQUEDA_WEB", "0")
    miembro = MiembroOpenAI()
    falsas = RespuestasOpenAIFalsas()
    miembro._cliente = NS(responses=falsas)
    miembro.preguntar("s", "m")
    assert "reasoning" not in falsas.peticiones[0] and miembro.descripcion() == "modelo-normal"


def test_extremo_claude_usa_esfuerzo_maximo_y_razonamiento(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "clave-falsa")
    peticiones = []

    class Flujo:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get_final_message(self):
            return NS(stop_reason="end_turn",
                      content=[NS(type="thinking", text="(razonamiento)"), NS(type="text", text="ok")])

    def stream(**parametros):
        peticiones.append(parametros)
        return Flujo()

    miembro = MiembroClaude(nivel="extremo")
    miembro._cliente = NS(messages=NS(stream=stream))
    assert miembro.preguntar("s", "m") == "ok"
    assert peticiones[0]["output_config"] == {"effort": "max"}
    assert peticiones[0]["thinking"] == {"type": "adaptive"}
    assert peticiones[0]["max_tokens"] == 64000


def test_error_no_relacionado_con_esfuerzo_se_propaga(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "clave-falsa")
    miembro = MiembroOpenAI(nivel="extremo")

    def falla(**_):
        raise ErrorPeticion("modelo inexistente")

    miembro._cliente = NS(responses=NS(create=falla))
    with pytest.raises(ErrorPeticion):
        miembro.preguntar("s", "m")  # agota la escala y propaga el error, sin bucle infinito


def test_consejo_extremo(tmp_path):
    a, b = MiembroFalso("Claude"), MiembroFalso("ChatGPT")
    sistemas = []
    original = a.preguntar
    a.preguntar = lambda s, m: (sistemas.append(s), original(s, m))[1]
    r = Consejo([a, b], registro=lambda *_: None, nivel="extremo").deliberar("T")
    assert r.nivel == "extremo" and all("NIVEL EXTREMO" in s for s in sistemas)
    md, _ = guardar(r, tmp_path)
    assert "**Nivel:** Extremo" in md.read_text(encoding="utf-8")
    with pytest.raises(ValueError):
        Consejo([a, b], nivel="maximo")


def test_web_valida_nivel(tmp_path):
    cliente = crear_app(tmp_path, fabricas=fabricas_falsas()).test_client()
    r = cliente.post("/api/consultas", data={"tarea": "T", "nivel": "maximo",
                                             "miembros": ["claude", "chatgpt"], "presidente": "claude"})
    assert r.status_code == 400


def test_cargar_env(tmp_path, monkeypatch):
    from council.config import cargar_env

    for var in ("CLAVE_A", "CLAVE_B", "CLAVE_VACIA", "CLAVE_EXISTENTE"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("CLAVE_EXISTENTE", "del-entorno")
    env = tmp_path / ".env"
    env.write_text('# comentario\nCLAVE_A=uno\nexport CLAVE_B="dos"\nCLAVE_VACIA=\n'
                   'CLAVE_EXISTENTE=del-archivo\nlinea sin igual\n', encoding="utf-8")
    assert sorted(cargar_env(env)) == ["CLAVE_A", "CLAVE_B"]
    import os
    assert os.environ["CLAVE_A"] == "uno" and os.environ["CLAVE_B"] == "dos"
    assert os.environ["CLAVE_EXISTENTE"] == "del-entorno"  # el entorno tiene prioridad
    assert cargar_env(tmp_path / "no-existe") == []


def test_verificar(monkeypatch):
    from council.verificar import verificar

    class Roto(MiembroFalso):
        def preguntar(self, sistema, mensaje):
            raise RuntimeError("modelo no encontrado")

    monkeypatch.setenv("CLAVE_OK", "x")
    monkeypatch.setenv("CLAVE_ROTA", "x")
    monkeypatch.delenv("CLAVE_FALTA", raising=False)
    fabricas = {
        "claude": lambda nivel: MiembroFalso("Claude"),
        "chatgpt": lambda nivel: Roto("ChatGPT"),
        "perplexity": lambda nivel: MiembroFalso("Perplexity"),
    }
    claves = {"claude": "CLAVE_OK", "chatgpt": "CLAVE_ROTA", "perplexity": "CLAVE_FALTA"}
    lineas = []
    assert verificar(["estandar"], fabricas, claves, escribir=lineas.append) == 1
    texto = "\n".join(lineas)
    assert "[ OK ] Claude" in texto and "[FALLO] ChatGPT" in texto and "modelo no encontrado" in texto
    assert "error inesperado" in texto
    assert "[ -- ] Perplexity: sin clave" in texto and "Hacen falta al menos dos" in texto

    monkeypatch.setenv("CLAVE_FALTA", "x")
    lineas.clear()
    assert verificar(["estandar", "extremo"], fabricas, claves, escribir=lineas.append) == 2
    assert "Consejo operativo con 2 miembros" in lineas[-1]


@pytest.mark.parametrize("estado,pista", [(401, "clave de API no es válida"), (404, "modelo indicado no existe"),
                                          (429, "límite de peticiones"), (400, "ha rechazado la petición")])
def test_explicar_errores(estado, pista):
    from council.verificar import _explicar

    class ErrorApi(Exception):
        status_code = estado

    assert pista in _explicar(ErrorApi("x"))


def test_principios_prohiben_ofrecimientos():
    from council import prompts

    assert "No hagas preguntas al usuario ni ofrezcas ampliaciones" in prompts.sistema()


def test_refinar_resultado():
    from council.consejo import refinar

    (a, _, _), consejo = consejo_de_tres()
    r = consejo.deliberar("Informe IRPF", "Doc", modo="informe_tributario")
    original = r.sintesis
    refinar(a, r, "Añade citas concretas", "Doc")
    assert r.sintesis == "## Respuesta final\nVersión corregida por Claude"
    assert r.versiones[0]["sintesis"] == original
    assert r.versiones[0]["instrucciones"] == "Añade citas concretas"
    peticion = a.llamadas[-1]
    # El presidente recibe todo el material anonimizado, su versión vigente y las indicaciones.
    assert "### Borrador de Miembro B" in peticion and original in peticion
    assert "Añade citas concretas" in peticion and "Doc" in peticion
    with pytest.raises(ValueError):
        refinar(a, r, "   ")


def test_servicio_refinar_seudonimiza_y_versiona(tmp_path):
    from council.servicio import refinar

    enviados = []

    class Espia(MiembroFalso):
        def preguntar(self, sistema, mensaje):
            enviados.append(mensaje)
            return super().preguntar(sistema, mensaje)

    fabricas = {"claude": lambda nivel: Espia("Claude"), "chatgpt": lambda nivel: Espia("ChatGPT")}
    salida = ejecutar("Cliente DNI 12345678Z", seudonimizar=True, carpeta=tmp_path,
                      registro=lambda *_: None, fabricas=fabricas)
    primero = salida.informe
    refinar(salida, "Usa también el correo ana@ejemplo.es y el DNI 12345678Z", registro=lambda *_: None,
            fabricas=fabricas)
    # Ni las indicaciones ni el resto del material salen con datos reales.
    assert "ana@ejemplo.es" not in enviados[-1] and "12345678Z" not in enviados[-1]
    assert "[DNI_1]" in enviados[-1] and "[EMAIL_1]" in enviados[-1]
    # La nueva versión se guarda aparte, con los datos restaurados y el historial.
    assert salida.informe != primero and salida.informe.name.endswith("_v2.md") and primero.exists()
    texto = salida.informe.read_text(encoding="utf-8")
    assert "Versión corregida por Claude" in texto and "Anexo III · Historial de versiones" in texto
    assert "ana@ejemplo.es" in texto and "**Versión:** 2" in texto
    assert texto.count("\n## Respuesta final") == 1 and "\n#### Respuesta final" in texto
    assert json.loads(salida.mapa.read_text(encoding="utf-8"))["[EMAIL_1]"] == "ana@ejemplo.es"


def test_web_refinar(tmp_path):
    cliente = crear_app(tmp_path, fabricas=fabricas_falsas()).test_client()

    def esperar(id_consulta):
        for _ in range(100):
            estado = cliente.get(f"/api/consultas/{id_consulta}").get_json()
            if estado["estado"] != "en_curso":
                return estado
            time.sleep(0.05)
        raise AssertionError("no terminó")

    r = cliente.post("/api/consultas", data={"tarea": "T", "miembros": ["claude", "chatgpt"],
                                             "presidente": "claude"})
    id_consulta = r.get_json()["id"]
    assert esperar(id_consulta)["version"] == 1
    assert cliente.post(f"/api/consultas/{id_consulta}/refinar", data={"instrucciones": ""}).status_code == 400
    assert cliente.post("/api/consultas/otra/refinar", data={"instrucciones": "x"}).status_code == 404
    r = cliente.post(f"/api/consultas/{id_consulta}/refinar", data={"instrucciones": "Más citas"})
    assert r.status_code == 202
    estado = esperar(id_consulta)
    assert estado["version"] == 2 and estado["error"] is None
    assert "Versión corregida por Claude" in estado["informe"]
    assert any("Versión 2 terminada" in m for m in estado["mensajes"])
    descarga = cliente.get(f"/api/consultas/{id_consulta}/informe")
    assert "_v2" in descarga.headers["Content-Disposition"]


def test_claude_busqueda_web_con_fuentes_y_pausa(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "clave-falsa")
    monkeypatch.setenv("BUSQUEDA_DOMINIOS", "boe.es, hacienda.gob.es")
    peticiones = []
    respuestas = [
        NS(stop_reason="pause_turn", content=[NS(type="text", text="Según el BOE, ", citations=None)]),
        NS(stop_reason="end_turn", content=[NS(type="text", text="el RDL 29/2026 regula…", citations=[
            NS(url="https://www.boe.es/buscar/act.php?id=BOE-A-2026-20823", title="RDL 29/2026")])]),
    ]

    class Flujo:
        def __init__(self, r):
            self.r = r

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get_final_message(self):
            return self.r

    def stream(**parametros):
        peticiones.append(parametros)
        return Flujo(respuestas[len(peticiones) - 1])

    miembro = MiembroClaude()
    miembro._cliente = NS(messages=NS(stream=stream))
    texto = miembro.preguntar("s", "m")
    assert texto.startswith("Según el BOE, el RDL 29/2026 regula…")
    assert "### Fuentes consultadas\n1. [RDL 29/2026](https://www.boe.es/buscar/act.php?id=BOE-A-2026-20823)" in texto
    assert "Claude" not in texto  # sin pistas de autoría para la revisión anónima
    herramienta = peticiones[0]["tools"][0]
    assert herramienta["name"] == "web_search" and herramienta["allowed_domains"] == ["boe.es", "hacienda.gob.es"]
    # Tras la pausa se reenvía lo generado para que el modelo continúe.
    assert peticiones[1]["messages"][-1]["role"] == "assistant"
    assert miembro.descripcion() == "claude-opus-5-5 · búsqueda web"


def test_openai_busqueda_rechazada_se_desactiva(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "clave-falsa")
    monkeypatch.delenv("OPENAI_MODEL", raising=False)
    monkeypatch.setenv("BUSQUEDA_DOMINIOS", "boe.es")
    peticiones = []

    def create(**parametros):
        peticiones.append(parametros)
        if "tools" in parametros:
            raise ErrorPeticion("web_search no admitido por este modelo")
        return NS(output_text="ok", output=[NS(type="message", content=[NS(annotations=[
            NS(type="url_citation", url="https://www.boe.es/x", title="BOE")])])])

    miembro = MiembroOpenAI()
    miembro._cliente = NS(responses=NS(create=create))
    assert peticiones == [] and miembro.busqueda
    texto = miembro.preguntar("s", "m")
    assert peticiones[0]["tools"] == [{"type": "web_search", "filters": {"allowed_domains": ["boe.es"]}}]
    assert "tools" not in peticiones[-1] and "[BOE](https://www.boe.es/x)" in texto
    assert miembro.descripcion() == "gpt-6.1-sol · sin búsqueda web (rechazada por el proveedor)"


def test_instrucciones_normas_recientes():
    from datetime import date

    from council import prompts

    texto = prompts.sistema()
    assert date.today().strftime("%d/%m/%Y") in texto
    assert "NUNCA afirmes que no existe" in texto


class ModeloNoEncontrado(Exception):
    status_code = 404


def test_cadena_de_modelos_de_reserva(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "clave-falsa")
    monkeypatch.setenv("OPENAI_MODEL", "gpt-futuro, gpt-6-sol")
    monkeypatch.setenv("BUSQUEDA_WEB", "1")
    modelos_pedidos = []

    def create(**parametros):
        modelos_pedidos.append(parametros["model"])
        if parametros["model"] == "gpt-futuro":
            raise ModeloNoEncontrado("The model `gpt-futuro` does not exist")
        return NS(output_text="ok", output=[])

    miembro = MiembroOpenAI()
    miembro._cliente = NS(responses=NS(create=create))
    assert miembro.preguntar("s", "m") == "ok" and miembro.preguntar("s", "m") == "ok"
    # Se descarta el modelo inexistente sin desactivar la búsqueda, y no se vuelve a pedir.
    assert modelos_pedidos == ["gpt-futuro", "gpt-6-sol", "gpt-6-sol"]
    assert miembro.busqueda and miembro.descripcion() == (
        "gpt-6-sol · búsqueda web · sustituye a gpt-futuro (no disponible)")


def test_cadena_agotada_propaga_error(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "clave-falsa")
    monkeypatch.setenv("OPENAI_MODEL", "solo-este")
    monkeypatch.setenv("BUSQUEDA_WEB", "0")

    def create(**_):
        raise ModeloNoEncontrado("model not found")

    miembro = MiembroOpenAI()
    miembro._cliente = NS(responses=NS(create=create))
    with pytest.raises(ModeloNoEncontrado):
        miembro.preguntar("s", "m")


def test_modelos_por_defecto(monkeypatch):
    from council.modelos import candidatos

    for var in ("CLAUDE_MODEL", "OPENAI_MODEL", "OPENAI_MODEL_EXTREMO"):
        monkeypatch.delenv(var, raising=False)
    assert candidatos("claude", "estandar") == ["claude-opus-5-5"]
    assert candidatos("chatgpt", "estandar")[0] == "gpt-6.1-sol"
    assert candidatos("chatgpt", "extremo")[0] == "gpt-6-astra"


def test_word_convierte_markdown(tmp_path):
    from docx import Document

    from council.consejo import Resultado
    from council.exportar import exportar

    r = Resultado(tarea="Consulta de [EMAIL_1]", modo="informe_tributario", presidente="Claude",
                  modelos={"Claude": "claude-opus-5-5"}, sintesis="""## Respuesta final
La ganancia **tributa** en la base del *ahorro* según el [art. 33 LIRPF](https://www.boe.es/x).
Continúa el mismo párrafo.

- Primer punto
  - Subpunto
1. Paso uno

| Concepto | Importe |
|---|---|
| Valor de transmisión | 300.000 € |

> Debe verificarse la vigencia.

## Nivel de confianza
Medio""")
    ruta = exportar(r, "docx", tmp_path / "p", restaurar=lambda t: t.replace("[EMAIL_1]", "a@b.es"))
    assert ruta.suffix == ".docx"
    doc = Document(ruta)
    parrafos = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    textos = [t for _, t in parrafos]
    assert ("Title", "Informe del consejo de LLM") in parrafos
    assert ("Heading 2", "Respuesta final") in parrafos and ("Heading 2", "Nivel de confianza") in parrafos
    cuerpo = next(p for p in doc.paragraphs if p.text.startswith("La ganancia"))
    assert cuerpo.text == ("La ganancia tributa en la base del ahorro según el art. 33 LIRPF "
                           "(https://www.boe.es/x). Continúa el mismo párrafo.")
    assert any(run.bold and run.text == "tributa" for run in cuerpo.runs)
    assert any(run.italic and run.text == "ahorro" for run in cuerpo.runs)
    assert ("List Bullet", "Primer punto") in parrafos and ("List Bullet 2", "Subpunto") in parrafos
    assert ("List Number", "Paso uno") in parrafos and ("Quote", "Debe verificarse la vigencia.") in parrafos
    metadatos, importes = doc.tables
    assert metadatos.cell(len(metadatos.rows) - 1, 1).text == "Consulta de a@b.es"
    assert [c.text for c in importes.rows[1].cells] == ["Valor de transmisión", "300.000 €"]
    assert not any("|---" in t for t in textos)


def test_servicio_y_web_generan_word(tmp_path):
    from docx import Document

    cliente = crear_app(tmp_path, fabricas=fabricas_falsas()).test_client()
    r = cliente.post("/api/consultas", data={"tarea": "Cliente DNI 12345678Z", "miembros": ["claude", "chatgpt"],
                                             "presidente": "claude", "seudonimizar": "on"})
    id_consulta = r.get_json()["id"]
    for _ in range(100):
        estado = cliente.get(f"/api/consultas/{id_consulta}").get_json()
        if estado["estado"] != "en_curso":
            break
        time.sleep(0.05)
    # El progreso informa de cada entrega.
    assert sum("✓" in m for m in estado["mensajes"]) == 5
    assert any("✓ Claude ha entregado la síntesis final" in m for m in estado["mensajes"])
    descarga = cliente.get(f"/api/consultas/{id_consulta}/word")
    assert descarga.status_code == 200 and "_presidente.docx" in descarga.headers["Content-Disposition"]
    ruta = tmp_path / "descargado.docx"
    ruta.write_bytes(descarga.data)
    texto = "\n".join(p.text for p in Document(ruta).paragraphs)
    tablas = "\n".join(c.text for t in Document(ruta).tables for f in t.rows for c in f.cells)
    assert "Síntesis de Claude" in texto and "12345678Z" in tablas  # datos restaurados en local
    assert "Borrador de" not in texto  # solo el informe del presidente, sin anexos
    assert cliente.get(f"/api/consultas/{id_consulta}/otro").status_code == 404


SINTESIS_EJEMPLO = """## Respuesta final
La ganancia **tributa** en la base del *ahorro* (art. 33 → 49 LIRPF) ✓ según el [BOE](https://www.boe.es/x).

1. Primer paso
2. Segundo paso

- Viñeta con `código`

| Concepto | Importe |
|---|---|
| Valor | 300.000 € |

> Verificar la vigencia.
"""


@pytest.mark.parametrize("formato", ["pdf", "html", "md", "txt"])
def test_exportar_otros_formatos(tmp_path, formato):
    from council.consejo import Resultado
    from council.exportar import exportar

    r = Resultado(tarea="Consulta de [DNI_1]", modo="informe_tributario", presidente="Claude",
                  modelos={"Claude": "claude-opus-5-5"}, sintesis=SINTESIS_EJEMPLO)
    ruta = exportar(r, formato, tmp_path / "informe", restaurar=lambda t: t.replace("[DNI_1]", "12345678Z"))
    assert ruta.suffix == f".{formato}" and ruta.stat().st_size > 0
    if formato == "pdf":
        from pypdf import PdfReader

        texto = "".join(p.extract_text() for p in PdfReader(ruta).pages)
        assert "Informe del consejo de LLM" in texto and "300.000 €" in texto and "12345678Z" in texto
        assert "Respuesta final" in texto and "Primer paso" in texto
        return
    texto = ruta.read_text(encoding="utf-8")
    assert "12345678Z" in texto and "300.000 €" in texto and "[DNI_1]" not in texto
    if formato == "html":
        assert "<strong>tributa</strong>" in texto and '<a href="https://www.boe.es/x">BOE</a>' in texto
        assert "<ol>" in texto and "<th>Concepto</th>" in texto and "<blockquote>" in texto
    elif formato == "md":
        assert SINTESIS_EJEMPLO.strip() in texto
    else:
        assert "**" not in texto and "|---" not in texto and "BOE (https://www.boe.es/x)" in texto
        assert "1. Primer paso" in texto and "2. Segundo paso" in texto


def test_pdf_sin_fuente_unicode_sustituye_simbolos(tmp_path, monkeypatch):
    from pypdf import PdfReader

    from council import exportar as modulo
    from council.consejo import Resultado

    monkeypatch.setattr(modulo, "_FUENTES", [])
    r = Resultado(tarea="T", modo="general", presidente="Claude", sintesis=SINTESIS_EJEMPLO)
    texto = "".join(p.extract_text() for p in PdfReader(modulo.exportar(r, "pdf", tmp_path / "x")).pages)
    assert "33 -> 49" in texto and "300.000 €" in texto


def test_web_descarga_en_todos_los_formatos(tmp_path):
    cliente = crear_app(tmp_path, fabricas=fabricas_falsas()).test_client()
    formatos = [f["id"] for f in cliente.get("/api/configuracion").get_json()["formatos"]]
    assert formatos == ["docx", "pdf", "html", "md", "txt"]
    r = cliente.post("/api/consultas", data={"tarea": "T", "miembros": ["claude", "chatgpt"], "presidente": "claude"})
    id_consulta = r.get_json()["id"]
    for _ in range(100):
        if cliente.get(f"/api/consultas/{id_consulta}").get_json()["estado"] != "en_curso":
            break
        time.sleep(0.05)
    for formato in formatos:
        descarga = cliente.get(f"/api/consultas/{id_consulta}/presidente/{formato}")
        assert descarga.status_code == 200, formato
        assert f"_presidente.{formato}" in descarga.headers["Content-Disposition"]
    assert cliente.get(f"/api/consultas/{id_consulta}/presidente/exe").status_code == 404


def _libro_excel(ruta):
    from datetime import datetime

    from openpyxl import Workbook

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Liquidación"
    hoja.append(["Concepto", "Importe", "Fecha"])
    hoja.append(["Valor de transmisión", 300000, datetime(2026, 3, 15)])
    hoja.append(["Valor de adquisición", 180000.5, datetime(2010, 6, 1)])
    hoja.append(["Ganancia", "=B2-B3", None])
    hoja.append(["Nota con | barra", True, None])
    hoja.append(["Tipo", 0.21, None])
    hoja["B6"].number_format = "0%"
    oculta = libro.create_sheet("Auxiliar")
    oculta.sheet_state = "hidden"
    oculta["C2"] = "dato oculto"
    libro.create_sheet("Vacía")
    libro.save(ruta)
    return ruta


def test_leer_excel(tmp_path):
    from council.documentos import leer_documento

    texto = leer_documento(_libro_excel(tmp_path / "liquidacion.xlsx"))
    assert texto.startswith("Libro de Excel con 3 hoja(s).")
    assert "### Hoja «Liquidación»" in texto and "| Fila | A | B | C |" in texto
    assert "| 2 | Valor de transmisión | 300000 | 15/03/2026 |" in texto
    assert "| 3 | Valor de adquisición | 180000.5 | 01/06/2010 |" in texto
    assert "| 5 | Nota con \\| barra | VERDADERO |  |" in texto
    assert "| 6 | Tipo | 21 % |  |" in texto
    # La fórmula se lista con su celda; openpyxl no calcula, así que se avisa de que falta el valor.
    assert "- B4: =B2-B3 → (sin valor calculado" in texto
    assert "### Hoja «Auxiliar» (hoja oculta)" in texto and "| 2 |  |  | dato oculto |" in texto
    assert "### Hoja «Vacía»\n(hoja vacía)" in texto


def test_leer_excel_recorta_hojas_grandes(tmp_path, monkeypatch):
    from openpyxl import Workbook

    from council import documentos

    monkeypatch.setattr(documentos, "MAX_FILAS", 5)
    libro = Workbook()
    for n in range(1, 21):
        libro.active.append([n, n * 2])
    libro.save(tmp_path / "grande.xlsx")
    texto = documentos.leer_documento(tmp_path / "grande.xlsx")
    assert "se muestran las primeras 5 de 20 filas" in texto and "| 5 | 5 | 10 |" in texto
    assert "\n| 6 | 6 | 12 |" not in texto


def test_leer_csv_espanol(tmp_path):
    from council.documentos import leer_documento

    ruta = tmp_path / "movimientos.csv"
    ruta.write_bytes("Fecha;Concepto;Importe\n01/02/2026;Cuota año;1.234,56\n".encode("cp1252"))
    texto = leer_documento(ruta)
    assert "| 2 | 01/02/2026 | Cuota año | 1.234,56 |" in texto


def test_leer_xls_antiguo(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    from council.documentos import leer_documento

    libro = xlwt.Workbook()
    hoja = libro.add_sheet("Datos")
    hoja.write(0, 0, "Base")
    hoja.write(0, 1, 1500.25)
    libro.save(str(tmp_path / "antiguo.xls"))
    texto = leer_documento(tmp_path / "antiguo.xls")
    assert "no permite leer las fórmulas" in texto and "| 1 | Base | 1500.25 |" in texto


def test_web_varios_documentos_con_excel(tmp_path):
    import io

    enviados = []

    class Espia(MiembroFalso):
        def preguntar(self, sistema, mensaje):
            enviados.append(mensaje)
            return super().preguntar(sistema, mensaje)

    fabricas = {"claude": lambda nivel="estandar": Espia("Claude"),
                "chatgpt": lambda nivel="estandar": Espia("ChatGPT")}
    cliente = crear_app(tmp_path, fabricas=fabricas).test_client()
    excel = _libro_excel(tmp_path / "liquidacion.xlsx").read_bytes()
    r = cliente.post("/api/consultas", content_type="multipart/form-data", data={
        "tarea": "Revisa la liquidación", "miembros": ["claude", "chatgpt"], "presidente": "claude",
        "documento": [(io.BytesIO(excel), "liquidacion.xlsx"),
                      (io.BytesIO("Contrato de compraventa…".encode()), "contrato.txt")],
    })
    assert r.status_code == 202, r.get_json()
    id_consulta = r.get_json()["id"]
    for _ in range(100):
        if cliente.get(f"/api/consultas/{id_consulta}").get_json()["estado"] != "en_curso":
            break
        time.sleep(0.05)
    primero = enviados[0]
    assert "===== DOCUMENTO 1: liquidacion.xlsx =====" in primero and "Valor de transmisión" in primero
    assert "===== DOCUMENTO 2: contrato.txt =====" in primero and "Contrato de compraventa" in primero
    rechazo = cliente.post("/api/consultas", content_type="multipart/form-data", data={
        "tarea": "T", "miembros": ["claude", "chatgpt"], "presidente": "claude",
        "documento": [(io.BytesIO(b"x"), "macro.exe")]})
    assert rechazo.status_code == 400 and "macro.exe" in rechazo.get_json()["error"]


def test_servidor_mcp_flujo_completo(tmp_path, monkeypatch):
    import anyio

    from council import mcp_servidor
    from council.trabajos import GestorConsultas

    monkeypatch.setattr(mcp_servidor, "ESPERA", 5)
    gestor = GestorConsultas(tmp_path, fabricas_falsas())
    servidor = mcp_servidor.crear_servidor(
        gestor, disponibles=lambda: {"claude": True, "chatgpt": True, "perplexity": False})

    def llamar(nombre, **argumentos):
        resultado = anyio.run(servidor.call_tool, nombre, argumentos)
        return resultado.content[0].text

    excel = _libro_excel(tmp_path / "liquidacion.xlsx")
    respuesta = llamar("convocar_consejo", tarea="Revisa la liquidación del cliente con DNI 12345678Z",
                       tipo_trabajo="informe_tributario", archivos=[str(excel)],
                       texto_adicional="Nota del usuario")
    assert "Consejo convocado con 2 documento(s)" in respuesta and "Claude, ChatGPT" in respuesta
    id_consulta = respuesta.split("id_consulta=")[1].split(".")[0]

    informe = llamar("estado_consulta", id_consulta=id_consulta)
    assert informe.startswith("Estado: terminada (versión 1).")
    assert "INFORME DEL PRESIDENTE:" in informe and "Síntesis de Claude" in informe
    assert "seudonimizados (1 marcadores" in informe and "12345678Z" not in informe
    assert "_presidente.docx" in informe

    assert "en preparación" in llamar("version_corregida", id_consulta=id_consulta, instrucciones="Más citas")
    assert "Versión corregida por Claude" in llamar("estado_consulta", id_consulta=id_consulta)

    ruta = llamar("exportar_informe", id_consulta=id_consulta, formato="pdf").split("guardado en: ")[1]
    assert ruta.endswith("_v2_presidente.pdf") and Path(ruta).is_file()

    assert "sin clave de API para perplexity" in llamar("convocar_consejo", tarea="T", miembros=["perplexity"])
    assert "formato no admitido" in llamar("convocar_consejo", tarea="T", archivos=["/tmp/x.exe"])
    assert "No existe la consulta" in llamar("exportar_informe", id_consulta="otra")


def test_instalar_en_claude_conserva_otros_servidores(tmp_path):
    from council.instalar_claude import registrar

    configuracion = tmp_path / "claude_desktop_config.json"
    configuracion.write_text(json.dumps({"mcpServers": {"otro": {"command": "x"}}, "tema": "oscuro"}))
    copia = registrar(configuracion, Path("/ruta/python"), Path("/ruta/consejo_mcp.py"))
    datos = json.loads(configuracion.read_text())
    assert datos["mcpServers"]["otro"] == {"command": "x"} and datos["tema"] == "oscuro"
    assert datos["mcpServers"]["consejo-llm"] == {"command": "/ruta/python", "args": ["/ruta/consejo_mcp.py"]}
    assert copia.is_file() and "otro" in copia.read_text()
    # Sin archivo previo, lo crea (y sin copia de seguridad).
    nueva = tmp_path / "nueva" / "claude_desktop_config.json"
    assert registrar(nueva, Path("/p"), Path("/s")) is None and "consejo-llm" in nueva.read_text()


# --- Auditoría (bucle 1) ------------------------------------------------------------------

class Fallido(MiembroFalso):
    def __init__(self, nombre, fase):
        super().__init__(nombre)
        self.fase = fase

    def preguntar(self, sistema, mensaje):
        fase = ("sintesis" if "PRESIDENTE" in mensaje else
                "revision" if "Respuesta de Miembro" in mensaje else "borrador")
        if fase == self.fase:
            raise RuntimeError(f"caída en {fase}")
        return super().preguntar(sistema, mensaje)


def test_consejo_sigue_si_falla_un_miembro():
    a, b, c = MiembroFalso("Claude"), Fallido("ChatGPT", "borrador"), MiembroFalso("Perplexity")
    r = Consejo([a, b, c], registro=lambda *_: None).deliberar("T")
    assert set(r.borradores) == {"Claude", "Perplexity"} and set(r.revisiones) == {"Claude", "Perplexity"}
    assert r.incidencias == ["ChatGPT no pudo completar su borrador: RuntimeError: caída en borrador"]
    # Las revisiones solo incluyen a los miembros activos, con alias coherentes con los materiales.
    revision = next(m for m in a.llamadas if "Respuesta de Miembro" in m)
    assert "### Respuesta de Miembro B" in revision and "Miembro C" not in revision
    assert "Síntesis de Claude" in r.sintesis


def test_consejo_cambia_de_presidente_si_falla_la_sintesis():
    a, b = Fallido("Claude", "sintesis"), MiembroFalso("ChatGPT")
    r = Consejo([a, b], registro=lambda *_: None).deliberar("T")
    assert r.presidente == "ChatGPT" and "Síntesis de ChatGPT" in r.sintesis
    assert "Claude no pudo completar su síntesis" in r.incidencias[0]


def test_consejo_falla_si_quedan_menos_de_dos():
    a, b = MiembroFalso("Claude"), Fallido("ChatGPT", "borrador")
    with pytest.raises(RuntimeError, match="Menos de dos miembros"):
        Consejo([a, b], registro=lambda *_: None).deliberar("T")


def test_incidencias_en_informe_y_exportacion(tmp_path):
    from council.exportar import exportar

    a, b, c = MiembroFalso("Claude"), Fallido("ChatGPT", "revision"), MiembroFalso("Perplexity")
    r = Consejo([a, b, c], registro=lambda *_: None).deliberar("T")
    md, _ = guardar(r, tmp_path)
    assert "**Incidencias:** ChatGPT no pudo completar su revisión" in md.read_text(encoding="utf-8")
    assert "ChatGPT no pudo completar su revisión" in exportar(r, "txt", tmp_path / "x").read_text(encoding="utf-8")


def test_respuestas_truncadas_se_avisan(monkeypatch):
    from council.providers import AVISO_TRUNCADO

    monkeypatch.setenv("ANTHROPIC_API_KEY", "x")
    monkeypatch.setenv("OPENAI_API_KEY", "x")
    monkeypatch.setenv("BUSQUEDA_WEB", "0")

    class Flujo:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get_final_message(self):
            return NS(stop_reason="max_tokens", content=[NS(type="text", text="Contrato…")])

    claude = MiembroClaude()
    claude._cliente = NS(messages=NS(stream=lambda **_: Flujo()))
    assert claude.preguntar("s", "m") == "Contrato…" + AVISO_TRUNCADO

    openai_ = MiembroOpenAI()
    openai_._cliente = NS(responses=NS(create=lambda **_: NS(output_text="Parcial", output=[], status="incomplete")))
    assert openai_.preguntar("s", "m").endswith(AVISO_TRUNCADO)


def test_pdf_escaneado_se_rechaza(tmp_path):
    from reportlab.pdfgen import canvas

    from council.documentos import leer_documento

    ruta = tmp_path / "notificacion.pdf"
    c = canvas.Canvas(str(ruta))
    c.rect(100, 100, 300, 300, fill=1)  # solo una «imagen», sin texto
    c.showPage()
    c.save()
    with pytest.raises(ValueError, match="escaneado"):
        leer_documento(ruta)


def test_word_incluye_tablas_en_orden(tmp_path):
    import docx

    from council.documentos import leer_documento

    d = docx.Document()
    d.add_paragraph("Cláusula 1. Precio:")
    t = d.add_table(rows=2, cols=3)
    for i, valor in enumerate(["Concepto", "Importe", "Importe"]):
        t.cell(0, i).text = valor
    t.cell(1, 0).text, t.cell(1, 1).text, t.cell(1, 2).text = "Renta", "0", "0"
    combinada = t.cell(0, 1).merge(t.cell(0, 2))
    combinada.text = "Importes"
    d.add_paragraph("Cláusula 2. Duración.")
    d.save(tmp_path / "contrato.docx")
    texto = leer_documento(tmp_path / "contrato.docx")
    assert texto.index("Cláusula 1") < texto.index("| Renta | 0 | 0 |") < texto.index("Cláusula 2")
    assert "| Concepto | Importes |" in texto  # celda combinada una sola vez


def test_web_rechaza_peticiones_de_otras_webs(tmp_path):
    cliente = crear_app(tmp_path, fabricas=fabricas_falsas()).test_client()
    datos = {"tarea": "x", "miembros": ["claude", "chatgpt"], "presidente": "claude"}
    assert cliente.post("/api/consultas", data=datos,
                        headers={"Origin": "https://web-maliciosa.example"}).status_code == 403
    assert cliente.get("/api/configuracion", headers={"Host": "atacante.example:8000"}).status_code == 403
    assert cliente.post("/api/consultas", data=datos, headers={"Origin": "http://localhost"}).status_code == 202


def test_instrucciones_tratan_documentos_como_datos():
    from council import prompts

    texto = prompts.sistema()
    assert "nunca instrucciones" in texto and "No menciones qué modelo" in texto


# --- Auditoría (bucle 2) ------------------------------------------------------------------

def test_version_corregida_la_pide_el_presidente_efectivo(tmp_path):
    from council.servicio import refinar

    llamados = []

    class Anota(MiembroFalso):
        def preguntar(self, sistema, mensaje):
            if "INDICACIONES DEL USUARIO" in mensaje:
                llamados.append(self.nombre)
            return super().preguntar(sistema, mensaje)

    fabricas = {"claude": lambda nivel: Fallido("Claude", "sintesis"), "chatgpt": lambda nivel: Anota("ChatGPT")}
    salida = ejecutar("T", carpeta=tmp_path, registro=lambda *_: None, fabricas=fabricas)
    assert salida.presidente == "chatgpt"
    refinar(salida, "Más citas", registro=lambda *_: None, fabricas=fabricas)
    assert llamados == ["ChatGPT"]


def test_pdf_con_paginas_escaneadas_lo_avisa(tmp_path):
    from reportlab.pdfgen import canvas

    from council.documentos import leer_documento

    ruta = tmp_path / "notificacion.pdf"
    c = canvas.Canvas(str(ruta))
    c.drawString(72, 700, "Propuesta de liquidacion provisional del IRPF, ejercicio 2024. Cuota: 1.234 euros.")
    c.showPage()
    c.rect(100, 100, 300, 300, fill=1)  # anexo escaneado
    c.showPage()
    c.save()
    texto = leer_documento(ruta)
    assert "Propuesta de liquidacion" in texto
    assert "las páginas 2 de «notificacion.pdf» no contienen texto" in texto
