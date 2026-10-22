"""Exportación del informe del presidente a varios formatos: Word, PDF, HTML, Markdown y texto.

Un único analizador convierte el Markdown que generan los modelos en bloques (títulos, párrafos,
listas, tablas, citas y código) con formato en línea (negrita, cursiva, código y enlaces); cada
formato tiene su propio generador a partir de esos bloques, de modo que todos muestran lo mismo.
"""

from __future__ import annotations

import html
import os
import re
import threading
from dataclasses import dataclass
from pathlib import Path

from .consejo import Resultado
from .prompts import MODOS

AVISO = ("Aviso: este informe ha sido generado por modelos de lenguaje y constituye un borrador de "
         "apoyo. No sustituye el asesoramiento de un profesional cualificado, que debe revisarlo "
         "antes de cualquier uso.")
PIE = "Borrador generado por el consejo de LLM · Requiere revisión profesional"
TITULO = "Informe del consejo de LLM"

FORMATOS = {  # clave: (extensión, nombre visible)
    "docx": (".docx", "Word"),
    "pdf": (".pdf", "PDF"),
    "html": (".html", "Página web (HTML)"),
    "md": (".md", "Markdown"),
    "txt": (".txt", "Texto sin formato"),
}


# --- Análisis del Markdown ----------------------------------------------------------------

@dataclass
class Segmento:
    texto: str
    estilo: str = ""  # "", "b", "i", "code" o "link"
    url: str = ""


@dataclass
class Bloque:
    tipo: str  # "titulo", "parrafo", "lista", "tabla", "cita", "codigo"
    texto: str = ""
    nivel: int = 1  # títulos: 1-4; listas: 1-2
    numerada: bool = False
    filas: list[list[str]] | None = None


# **negrita**, *cursiva* o _cursiva_, `código` y [texto](url)
_EN_LINEA = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|(?<!\w)_(.+?)_(?!\w)|`(.+?)`|\[([^\]]+)\]\(([^)]+)\)")
_LISTA = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_INICIO_BLOQUE = re.compile(r"^\s*(#|>|\||```|[-*+]\s|\d+[.)]\s)")


def segmentos(texto: str) -> list[Segmento]:
    resultado, posicion = [], 0
    for m in _EN_LINEA.finditer(texto):
        if m.start() > posicion:
            resultado.append(Segmento(texto[posicion:m.start()]))
        if m.group(1) is not None:
            resultado.append(Segmento(m.group(1), "b"))
        elif m.group(2) is not None or m.group(3) is not None:
            resultado.append(Segmento(m.group(2) or m.group(3), "i"))
        elif m.group(4) is not None:
            resultado.append(Segmento(m.group(4), "code"))
        else:
            resultado.append(Segmento(m.group(5), "link", m.group(6)))
        posicion = m.end()
    if posicion < len(texto):
        resultado.append(Segmento(texto[posicion:]))
    return resultado


def texto_plano(texto: str) -> str:
    return "".join(s.texto + (f" ({s.url})" if s.url else "") for s in segmentos(texto))


def _celdas(linea: str) -> list[str]:
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def bloques(texto: str) -> list[Bloque]:
    lineas, resultado, i = texto.splitlines(), [], 0
    while i < len(lineas):
        linea = lineas[i]
        limpia = linea.strip()
        if not limpia or re.fullmatch(r"[-*_]{3,}", limpia):
            i += 1
        elif limpia.startswith("```"):
            codigo, i = [], i + 1
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                codigo.append(lineas[i])
                i += 1
            resultado.append(Bloque("codigo", "\n".join(codigo)))
            i += 1
        elif limpia.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                if not re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", lineas[i]):  # fila separadora
                    filas.append(_celdas(lineas[i]))
                i += 1
            if filas:
                ancho = max(len(f) for f in filas)
                resultado.append(Bloque("tabla", filas=[f + [""] * (ancho - len(f)) for f in filas]))
        elif m := re.match(r"^(#{1,6})\s+(.*)$", limpia):
            resultado.append(Bloque("titulo", m.group(2).strip("# "), nivel=min(len(m.group(1)), 4)))
            i += 1
        elif limpia.startswith(">"):
            resultado.append(Bloque("cita", limpia.lstrip("> ").strip()))
            i += 1
        elif m := _LISTA.match(linea):
            nivel = 2 if len(m.group(1).replace("\t", "    ")) >= 2 else 1
            resultado.append(Bloque("lista", m.group(3), nivel=nivel, numerada=m.group(2)[0].isdigit()))
            i += 1
        else:
            # Las líneas consecutivas forman un mismo párrafo, como en Markdown.
            partes, i = [limpia], i + 1
            while i < len(lineas) and lineas[i].strip() and not _INICIO_BLOQUE.match(lineas[i]):
                partes.append(lineas[i].strip())
                i += 1
            resultado.append(Bloque("parrafo", " ".join(partes)))
    return resultado


def metadatos(r: Resultado, restaurar) -> list[tuple[str, str]]:
    datos = [
        ("Fecha (UTC)", r.fecha[:19].replace("T", " ")),
        ("Tipo de trabajo", MODOS[r.modo][0]),
        ("Nivel", "Extremo (modelos más capaces y esfuerzo máximo)" if r.nivel == "extremo" else "Estándar"),
        ("Miembros", ", ".join(f"{n} ({m})" for n, m in r.modelos.items())),
        ("Presidente", r.presidente),
    ]
    if r.versiones:
        datos.append(("Versión", f"{len(r.versiones) + 1} (corregida a petición del usuario)"))
    if r.incidencias:
        datos.append(("Incidencias", "; ".join(r.incidencias)))
    datos.append(("Tarea", restaurar(r.tarea)))
    return datos


# --- Generadores por formato ---------------------------------------------------------------

def _word(r: Resultado, sintesis: str, datos, ruta: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    def en_linea(parrafo, texto: str) -> None:
        for s in segmentos(texto):
            if s.estilo == "link":
                parrafo.add_run(s.texto)
                parrafo.add_run(f" ({s.url})").font.color.rgb = RGBColor(0x2F, 0x4F, 0x8F)
                continue
            run = parrafo.add_run(s.texto)
            run.bold = s.estilo == "b" or None
            run.italic = s.estilo == "i" or None
            if s.estilo == "code":
                run.font.name = "Courier New"

    from docx.shared import Cm

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(TITULO, level=0)
    run = doc.add_paragraph().add_run(AVISO)
    run.italic, run.font.size = True, Pt(9)

    tabla = doc.add_table(rows=len(datos), cols=2)
    tabla.style = "Table Grid"
    for i, (etiqueta, valor) in enumerate(datos):
        tabla.cell(i, 0).text = etiqueta
        tabla.cell(i, 0).paragraphs[0].runs[0].bold = True
        tabla.cell(i, 1).text = valor
        tabla.cell(i, 0).width, tabla.cell(i, 1).width = Cm(4), Cm(12.5)
    tabla.autofit = False
    tabla.columns[0].width, tabla.columns[1].width = Cm(4), Cm(12.5)
    doc.add_paragraph()

    for b in bloques(sintesis):
        if b.tipo == "titulo":
            en_linea(doc.add_heading(level=b.nivel), b.texto)
        elif b.tipo == "parrafo":
            en_linea(doc.add_paragraph(), b.texto)
        elif b.tipo == "lista":
            estilo = ("List Number" if b.numerada else "List Bullet") + (" 2" if b.nivel == 2 else "")
            en_linea(doc.add_paragraph(style=estilo), b.texto)
        elif b.tipo == "cita":
            en_linea(doc.add_paragraph(style="Quote"), b.texto)
        elif b.tipo == "codigo":
            for linea in b.texto.splitlines():
                run = doc.add_paragraph().add_run(linea)
                run.font.name, run.font.size = "Courier New", Pt(9)
        elif b.tipo == "tabla":
            t = doc.add_table(rows=len(b.filas), cols=len(b.filas[0]))
            t.style = "Table Grid"
            for i, fila in enumerate(b.filas):
                for j, celda in enumerate(fila):
                    parrafo = t.cell(i, j).paragraphs[0]
                    en_linea(parrafo, celda)
                    if i == 0:
                        for run in parrafo.runs:
                            run.bold = True
            doc.add_paragraph()

    pie = doc.sections[0].footer.paragraphs[0]
    pie.text = PIE
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(ruta)


# Fuentes TrueType con acentos, «€», flechas, etc. (normal, negrita, cursiva, negrita cursiva).
_FUENTES = [
    ["/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
     "/System/Library/Fonts/Supplemental/Arial Italic.ttf", "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf"],
    ["/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf",
     "/Library/Fonts/Arial Italic.ttf", "/Library/Fonts/Arial Bold Italic.ttf"],
    ["C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf",
     "C:/Windows/Fonts/ariali.ttf", "C:/Windows/Fonts/arialbi.ttf"],
    ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"],
]
# Sustituciones si no hay fuente TrueType: Helvetica solo cubre el juego de caracteres de Windows.
_EQUIVALENCIAS = {"→": "->", "←": "<-", "⇒": "=>", "≥": ">=", "≤": "<=", "≠": "!=", "✓": "OK",
                  "✔": "OK", "✗": "X", "✅": "OK", "⚠️": "(!)", "⚠": "(!)", "•": "-", "≈": "~"}


def _fuente_pdf() -> tuple[str, bool]:
    """Registra la primera fuente TrueType disponible; si no hay, usa Helvetica."""
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for rutas in _FUENTES:
        if not Path(rutas[0]).is_file():
            continue
        try:
            nombres = ["Informe", "Informe-Negrita", "Informe-Cursiva", "Informe-NegritaCursiva"]
            for nombre, ruta in zip(nombres, rutas):
                pdfmetrics.registerFont(TTFont(nombre, ruta if Path(ruta).is_file() else rutas[0]))
            for (negrita, cursiva), nombre in zip([(0, 0), (1, 0), (0, 1), (1, 1)], nombres):
                addMapping("Informe", negrita, cursiva, nombre)
            return "Informe", True
        except Exception:
            continue
    return "Helvetica", False


def _pdf(r: Resultado, sintesis: str, datos, ruta: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (ListFlowable, ListItem, Paragraph, Preformatted, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    fuente, unicode_completo = _fuente_pdf()

    def limpiar(texto: str) -> str:
        if not unicode_completo:
            for original, equivalente in _EQUIVALENCIAS.items():
                texto = texto.replace(original, equivalente)
        return texto

    def en_linea(texto: str) -> str:
        partes = []
        for s in segmentos(limpiar(texto)):
            t = html.escape(s.texto, quote=False)
            if s.estilo == "b":
                t = f"<b>{t}</b>"
            elif s.estilo == "i":
                t = f"<i>{t}</i>"
            elif s.estilo == "code":
                t = f'<font face="Courier">{t}</font>'
            elif s.estilo == "link":
                url = html.escape(s.url)
                t = f'{t} (<link href="{url}" color="#2F4F8F">{url}</link>)'
            partes.append(t)
        return "".join(partes)

    base = getSampleStyleSheet()
    normal = ParagraphStyle("normal", parent=base["Normal"], fontName=fuente, fontSize=10.5, leading=14,
                            spaceAfter=6)
    # La negrita de los títulos se aplica con <b>, para que use la variante de la fuente registrada.
    titulos = {n: ParagraphStyle(f"t{n}", parent=base[f"Heading{n}"], fontName=fuente,
                                 textColor=colors.HexColor("#1D2B4F")) for n in range(1, 5)}
    titulo = ParagraphStyle("titulo", parent=base["Title"], fontName=fuente)
    aviso = ParagraphStyle("aviso", parent=normal, fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    cita = ParagraphStyle("cita", parent=normal, leftIndent=12, borderPadding=6, spaceBefore=8, spaceAfter=10,
                          backColor=colors.HexColor("#FFF6E0"), borderColor=colors.HexColor("#E6C46A"),
                          borderWidth=0.5)
    celda = ParagraphStyle("celda", parent=normal, fontSize=9.5, leading=12, spaceAfter=0)

    def tabla(filas: list[list[str]], cabecera: bool) -> Table:
        contenido = [[Paragraph(en_linea(c) if not (cabecera and i == 0) else f"<b>{en_linea(c)}</b>", celda)
                      for c in fila] for i, fila in enumerate(filas)]
        t = Table(contenido, repeatRows=1 if cabecera else 0, hAlign="LEFT",
                  colWidths=[(A4[0] - 4 * cm) / len(filas[0])] * len(filas[0]))
        estilo = [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BBBBBB")),
                  ("VALIGN", (0, 0), (-1, -1), "TOP")]
        if cabecera:
            estilo.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF9")))
        t.setStyle(TableStyle(estilo))
        return t

    historia = [Paragraph(f"<b>{TITULO}</b>", titulo), Paragraph(html.escape(AVISO), aviso), Spacer(1, 6)]
    meta = tabla([[f"**{e}**", v] for e, v in datos], cabecera=False)
    meta._argW = [4 * cm, A4[0] - 8 * cm]
    historia += [meta, Spacer(1, 12)]

    lista_actual: list = []
    numero = 0

    def cerrar_lista() -> None:
        nonlocal numero
        if lista_actual:
            historia.append(ListFlowable(list(lista_actual), bulletType="bullet", start="•",
                                         leftIndent=14, bulletFontName=fuente, bulletFontSize=10.5))
            lista_actual.clear()
        numero = 0

    for b in bloques(sintesis):
        if b.tipo != "lista":
            cerrar_lista()
        if b.tipo == "titulo":
            historia.append(Paragraph(f"<b>{en_linea(b.texto)}</b>", titulos[b.nivel]))
        elif b.tipo == "parrafo":
            historia.append(Paragraph(en_linea(b.texto), normal))
        elif b.tipo == "lista":
            if b.numerada:
                numero += 1
                marca = f"{numero}."
            else:
                marca = "–" if b.nivel == 2 else "•"
            lista_actual.append(ListItem(Paragraph(en_linea(b.texto), normal), value=marca,
                                         leftIndent=24 if b.nivel == 2 else 14))
        elif b.tipo == "cita":
            historia.append(Paragraph(en_linea(b.texto), cita))
        elif b.tipo == "codigo":
            historia.append(Preformatted(limpiar(b.texto), ParagraphStyle("codigo", fontName="Courier", fontSize=8.5)))
        elif b.tipo == "tabla":
            historia += [tabla(b.filas, cabecera=True), Spacer(1, 8)]
    cerrar_lista()

    def pie(canvas, documento) -> None:
        canvas.saveState()
        canvas.setFont(fuente, 8)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawCentredString(A4[0] / 2, 1.2 * cm, f"{limpiar(PIE)} · Página {documento.page}")
        canvas.restoreState()

    SimpleDocTemplate(str(ruta), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title=TITULO, author="Consejo de LLM").build(
        historia, onFirstPage=pie, onLaterPages=pie)


def _html(r: Resultado, sintesis: str, datos, ruta: Path) -> None:
    def en_linea(texto: str) -> str:
        partes = []
        for s in segmentos(texto):
            t = html.escape(s.texto)
            partes.append({"b": f"<strong>{t}</strong>", "i": f"<em>{t}</em>", "code": f"<code>{t}</code>",
                           "link": f'<a href="{html.escape(s.url)}">{t}</a>'}.get(s.estilo, t))
        return "".join(partes)

    cuerpo, lista_abierta = [], None
    for b in bloques(sintesis):
        etiqueta = ("ol" if b.numerada else "ul") if b.tipo == "lista" else None
        if lista_abierta and etiqueta != lista_abierta:
            cuerpo.append(f"</{lista_abierta}>")
            lista_abierta = None
        if b.tipo == "titulo":
            cuerpo.append(f"<h{b.nivel + 1}>{en_linea(b.texto)}</h{b.nivel + 1}>")
        elif b.tipo == "parrafo":
            cuerpo.append(f"<p>{en_linea(b.texto)}</p>")
        elif b.tipo == "lista":
            if not lista_abierta:
                cuerpo.append(f"<{etiqueta}>")
                lista_abierta = etiqueta
            sangria = ' class="n2"' if b.nivel == 2 else ""
            cuerpo.append(f"<li{sangria}>{en_linea(b.texto)}</li>")
        elif b.tipo == "cita":
            cuerpo.append(f"<blockquote>{en_linea(b.texto)}</blockquote>")
        elif b.tipo == "codigo":
            cuerpo.append(f"<pre>{html.escape(b.texto)}</pre>")
        elif b.tipo == "tabla":
            filas = [f"<tr>{''.join(f'<th>{en_linea(c)}</th>' for c in b.filas[0])}</tr>"]
            filas += [f"<tr>{''.join(f'<td>{en_linea(c)}</td>' for c in f)}</tr>" for f in b.filas[1:]]
            cuerpo.append(f"<div class=\"tabla\"><table>{''.join(filas)}</table></div>")
    if lista_abierta:
        cuerpo.append(f"</{lista_abierta}>")

    meta = "".join(f"<tr><th>{html.escape(e)}</th><td>{html.escape(v)}</td></tr>" for e, v in datos)
    ruta.write_text(f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{TITULO}</title>
<style>
  body {{ font: 16px/1.6 -apple-system, "Segoe UI", Calibri, sans-serif; color: #1d1d1f; max-width: 820px;
         margin: 0 auto; padding: 32px 16px; }}
  h1 {{ margin-bottom: 4px; }} h2, h3, h4, h5 {{ color: #1d2b4f; }}
  .aviso {{ font-size: .85rem; color: #555; font-style: italic; }}
  table {{ border-collapse: collapse; margin: 12px 0; }} .tabla {{ overflow-x: auto; }}
  th, td {{ border: 1px solid #bbb; padding: 6px 10px; text-align: left; vertical-align: top; }}
  .tabla th {{ background: #e8eef9; }} .meta th {{ background: #f4f4f4; white-space: nowrap; }}
  blockquote {{ margin: 0; padding: 8px 14px; border-left: 4px solid #e6c46a; background: #fff6e0; }}
  pre, code {{ font-family: Menlo, Consolas, monospace; font-size: .9em; }} pre {{ white-space: pre-wrap; }}
  li.n2 {{ margin-left: 24px; list-style-type: circle; }} a {{ color: #2f4f8f; }}
  footer {{ margin-top: 40px; font-size: .8rem; color: #777; text-align: center; }}
</style></head><body>
<h1>{TITULO}</h1>
<p class="aviso">{html.escape(AVISO)}</p>
<table class="meta">{meta}</table>
{chr(10).join(cuerpo)}
<footer>{html.escape(PIE)}</footer>
</body></html>
""", encoding="utf-8")


def _markdown(r: Resultado, sintesis: str, datos, ruta: Path) -> None:
    meta = "\n".join(f"- **{e}:** {v}" for e, v in datos)
    ruta.write_text(f"# {TITULO}\n\n> {AVISO}\n\n{meta}\n\n{sintesis}\n\n---\n_{PIE}_\n", encoding="utf-8")


def _texto(r: Resultado, sintesis: str, datos, ruta: Path) -> None:
    salida = [TITULO.upper(), "=" * len(TITULO), "", AVISO, ""]
    salida += [f"{e}: {v}" for e, v in datos] + [""]
    numero = 0
    for b in bloques(sintesis):
        numero = numero + 1 if (b.tipo == "lista" and b.numerada) else 0
        if b.tipo == "titulo":
            t = texto_plano(b.texto)
            salida += ["", t.upper() if b.nivel <= 2 else t, ("-" * len(t)) if b.nivel <= 2 else ""]
        elif b.tipo == "parrafo":
            salida += [texto_plano(b.texto), ""]
        elif b.tipo == "lista":
            marca = f"{numero}." if b.numerada else ("-" if b.nivel == 1 else "  ·")
            salida.append(f"{'  ' * (b.nivel - 1)}{marca} {texto_plano(b.texto)}")
        elif b.tipo == "cita":
            salida += [f"  » {texto_plano(b.texto)}", ""]
        elif b.tipo == "codigo":
            salida += [f"    {l}" for l in b.texto.splitlines()] + [""]
        elif b.tipo == "tabla":
            salida += [" | ".join(texto_plano(c) for c in f) for f in b.filas] + [""]
    salida += ["", PIE]
    ruta.write_text("\n".join(salida) + "\n", encoding="utf-8")


_GENERADORES = {"docx": _word, "pdf": _pdf, "html": _html, "md": _markdown, "txt": _texto}


def exportar(r: Resultado, formato: str, ruta: str | Path, restaurar=None) -> Path:
    """Guarda el informe del presidente (sin anexos) en el formato indicado y devuelve su ruta."""
    if formato not in _GENERADORES:
        raise ValueError(f"Formato no admitido: {formato}. Opciones: {', '.join(FORMATOS)}")
    restaurar = restaurar or (lambda t: t)
    ruta = Path(ruta).with_suffix(FORMATOS[formato][0])
    # Se escribe en un archivo temporal y se renombra: así nunca se lee un archivo a medio escribir.
    temporal = ruta.with_name(f".{ruta.stem}.{os.getpid()}.{threading.get_ident()}{ruta.suffix}")
    try:
        _GENERADORES[formato](r, restaurar(r.sintesis), metadatos(r, restaurar), temporal)
        os.replace(temporal, ruta)
    finally:
        temporal.unlink(missing_ok=True)
    return ruta
