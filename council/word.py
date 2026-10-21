"""Exportación del informe del presidente a Word (.docx).

Convierte el Markdown que generan los modelos (títulos, listas, tablas, citas, negritas,
cursivas, código y enlaces) en un documento con estilos reales de Word, editable.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .consejo import Resultado
from .prompts import MODOS

AVISO = ("Aviso: este informe ha sido generado por modelos de lenguaje y constituye un borrador de "
         "apoyo. No sustituye el asesoramiento de un profesional cualificado, que debe revisarlo "
         "antes de cualquier uso.")

# **negrita**, *cursiva* o _cursiva_, `código` y [texto](url)
_EN_LINEA = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|(?<!\w)_(.+?)_(?!\w)|`(.+?)`|\[([^\]]+)\]\(([^)]+)\))")
_LISTA = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")


def _texto_en_linea(parrafo, texto: str) -> None:
    """Añade `texto` al párrafo interpretando el formato en línea de Markdown."""
    posicion = 0
    for m in _EN_LINEA.finditer(texto):
        if m.start() > posicion:
            parrafo.add_run(texto[posicion:m.start()])
        if m.group(2) is not None:
            parrafo.add_run(m.group(2)).bold = True
        elif m.group(3) is not None or m.group(4) is not None:
            parrafo.add_run(m.group(3) or m.group(4)).italic = True
        elif m.group(5) is not None:
            run = parrafo.add_run(m.group(5))
            run.font.name = "Courier New"
        else:
            parrafo.add_run(m.group(6))
            enlace = parrafo.add_run(f" ({m.group(7)})")
            enlace.font.color.rgb = RGBColor(0x2F, 0x4F, 0x8F)
        posicion = m.end()
    if posicion < len(texto):
        parrafo.add_run(texto[posicion:])


def _celdas(linea: str) -> list[str]:
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def _tabla(doc, filas: list[str]) -> None:
    datos = [_celdas(f) for f in filas if not re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", f)]
    if not datos:
        return
    columnas = max(len(f) for f in datos)
    tabla = doc.add_table(rows=len(datos), cols=columnas)
    tabla.style = "Table Grid"
    for i, fila in enumerate(datos):
        for j in range(columnas):
            celda = tabla.cell(i, j)
            celda.text = ""
            _texto_en_linea(celda.paragraphs[0], fila[j] if j < len(fila) else "")
            if i == 0:
                for run in celda.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


def anadir_markdown(doc, texto: str) -> None:
    """Añade al documento el contenido Markdown indicado."""
    lineas = texto.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        limpia = linea.strip()
        if not limpia or re.fullmatch(r"[-*_]{3,}", limpia):
            i += 1
        elif limpia.startswith("```"):
            i += 1
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                run = doc.add_paragraph().add_run(lineas[i])
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                i += 1
            i += 1
        elif limpia.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(lineas[i])
                i += 1
            _tabla(doc, filas)
        elif m := re.match(r"^(#{1,6})\s+(.*)$", limpia):
            titulo = doc.add_heading(level=min(len(m.group(1)), 4))
            _texto_en_linea(titulo, m.group(2).strip("# "))
            i += 1
        elif limpia.startswith(">"):
            _texto_en_linea(doc.add_paragraph(style="Quote"), limpia.lstrip("> ").strip())
            i += 1
        elif m := _LISTA.match(linea):
            numerada = m.group(2)[0].isdigit()
            nivel = 2 if len(m.group(1).replace("\t", "    ")) >= 2 else 1
            estilo = ("List Number" if numerada else "List Bullet") + (" 2" if nivel == 2 else "")
            _texto_en_linea(doc.add_paragraph(style=estilo), m.group(3))
            i += 1
        else:
            # Las líneas consecutivas forman un mismo párrafo, como en Markdown.
            partes = [limpia]
            i += 1
            while (i < len(lineas) and lineas[i].strip()
                   and not re.match(r"^\s*(#|>|\||```|[-*+]\s|\d+[.)]\s)", lineas[i])):
                partes.append(lineas[i].strip())
                i += 1
            _texto_en_linea(doc.add_paragraph(), " ".join(partes))


def informe_presidente(r: Resultado, ruta: str | Path, restaurar=None) -> Path:
    """Guarda en Word la respuesta del presidente (sin los anexos de borradores y revisiones)."""
    restaurar = restaurar or (lambda t: t)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Informe del consejo de LLM", level=0)
    aviso = doc.add_paragraph()
    run = aviso.add_run(AVISO)
    run.italic = True
    run.font.size = Pt(9)

    datos = [
        ("Fecha (UTC)", r.fecha[:19].replace("T", " ")),
        ("Tipo de trabajo", MODOS[r.modo][0]),
        ("Nivel", "Extremo (modelos más capaces y esfuerzo máximo)" if r.nivel == "extremo" else "Estándar"),
        ("Miembros", ", ".join(f"{n} ({m})" for n, m in r.modelos.items())),
        ("Presidente", r.presidente),
    ]
    if r.versiones:
        datos.append(("Versión", f"{len(r.versiones) + 1} (corregida a petición del usuario)"))
    datos.append(("Tarea", restaurar(r.tarea)))
    tabla = doc.add_table(rows=len(datos), cols=2)
    tabla.style = "Table Grid"
    for i, (etiqueta, valor) in enumerate(datos):
        tabla.cell(i, 0).text = etiqueta
        tabla.cell(i, 0).paragraphs[0].runs[0].bold = True
        tabla.cell(i, 1).text = valor
    doc.add_paragraph()

    anadir_markdown(doc, restaurar(r.sintesis))

    pie = doc.sections[0].footer.paragraphs[0]
    pie.text = "Borrador generado por el consejo de LLM · Requiere revisión profesional"
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER

    ruta = Path(ruta)
    doc.save(ruta)
    return ruta
