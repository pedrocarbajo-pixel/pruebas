"""Lectura de los documentos que se adjuntan a una consulta.

Formatos: texto (.txt, .md), Word (.docx), PDF y hojas de cálculo (.xlsx, .xlsm, .xls, .csv).
Las hojas de cálculo se convierten en tablas con letras de columna y números de fila, para que
los modelos puedan citar celdas concretas («B12»), y se listan las fórmulas para poder revisar
los cálculos.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time
from pathlib import Path

EXTENSIONES = {".txt", ".md", ".docx", ".pdf", ".xlsx", ".xlsm", ".xls", ".csv"}

MAX_FILAS = 2000  # por hoja
MAX_COLUMNAS = 50
MAX_FORMULAS = 300  # por hoja


def leer_documento(ruta: str | Path) -> str:
    ruta = Path(ruta)
    sufijo = ruta.suffix.lower()
    if sufijo == ".docx":
        return _word(ruta)
    if sufijo == ".pdf":
        return _pdf(ruta)
    if sufijo in (".xlsx", ".xlsm"):
        return _excel(ruta)
    if sufijo == ".xls":
        return _excel_antiguo(ruta)
    if sufijo == ".csv":
        return _csv(ruta)
    return ruta.read_text(encoding="utf-8")


def combinar(documentos: list[tuple[str, str]]) -> str | None:
    """Une varios documentos (nombre, texto) en un solo bloque, identificando cada uno."""
    if not documentos:
        return None
    return "\n\n".join(f"===== DOCUMENTO {i}: {nombre} =====\n{texto}"
                       for i, (nombre, texto) in enumerate(documentos, 1))


def leer_documentos(rutas: list[str | Path]) -> str | None:
    return combinar([(Path(r).name, leer_documento(r)) for r in rutas])


def _word(ruta: Path) -> str:
    """Texto de un Word en el orden del documento, incluidas las tablas (frecuentes en contratos)."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    documento = docx.Document(ruta)
    partes = []
    for elemento in documento.element.body.iterchildren():
        etiqueta = elemento.tag.rsplit("}", 1)[-1]
        if etiqueta == "p":
            partes.append(Paragraph(elemento, documento).text)
        elif etiqueta == "tbl":
            filas = []
            for fila in Table(elemento, documento).rows:
                celdas, anterior = [], None
                for celda in fila.cells:
                    # Una celda combinada aparece repetida (mismo elemento): se muestra una vez.
                    if celda._tc is not anterior:
                        celdas.append(celda.text.replace("|", "\\|").replace("\n", " ").strip())
                    anterior = celda._tc
                filas.append("| " + " | ".join(celdas) + " |")
            partes.append("\n".join(filas))
    return "\n".join(partes)


def _pdf(ruta: Path) -> str:
    from pypdf import PdfReader

    paginas = PdfReader(ruta).pages
    texto = "\n".join(pagina.extract_text() or "" for pagina in paginas)
    # Un PDF escaneado (una imagen) no contiene texto: analizarlo «vacío» sería engañoso.
    if len(texto.strip()) < 20 * max(len(paginas), 1):
        raise ValueError(f"«{ruta.name}» parece un PDF escaneado (imagen sin texto). Aplíquele "
                         "reconocimiento de texto (OCR), por ejemplo con Vista Previa o Adobe Acrobat, "
                         "y vuelva a adjuntarlo.")
    return texto


# --- Hojas de cálculo ------------------------------------------------------------------------

def _letra(columna: int) -> str:
    """1 → A, 27 → AA."""
    letras = ""
    while columna:
        columna, resto = divmod(columna - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


def _valor(v, formato: str = "") -> str:
    if v is None:
        return ""
    if "%" in (formato or "") and isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{round(v * 100, 8):g} %"  # como lo muestra Excel: 0,21 → 21 %
    if isinstance(v, bool):
        return "VERDADERO" if v else "FALSO"
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M") if v.time() != time(0) else v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(round(v, 10))
    return str(v).replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()


def _tabla(nombre: str, filas: list[list], formulas: list[tuple[str, str, str]] | None = None,
           nota: str = "") -> str:
    """Convierte una hoja (valores ya formateados como texto) en tabla Markdown con letras de
    columna y números de fila."""
    # Se descartan las filas y columnas vacías del final.
    filas = [[v if isinstance(v, str) else _valor(v) for v in fila] for fila in filas]
    while filas and not any(filas[-1]):
        filas.pop()
    ancho = max((max((i + 1 for i, v in enumerate(f) if v), default=0) for f in filas), default=0)
    cabecera = f"### Hoja «{nombre}»{nota}"
    if not ancho:
        return f"{cabecera}\n(hoja vacía)"
    avisos = []
    if ancho > MAX_COLUMNAS:
        avisos.append(f"se muestran las primeras {MAX_COLUMNAS} de {ancho} columnas")
        ancho = MAX_COLUMNAS
    if len(filas) > MAX_FILAS:
        avisos.append(f"se muestran las primeras {MAX_FILAS} de {len(filas)} filas")
        filas = filas[:MAX_FILAS]
    lineas = [cabecera]
    if avisos:
        lineas.append(f"(Hoja recortada: {'; '.join(avisos)}.)")
    lineas.append("| Fila | " + " | ".join(_letra(c) for c in range(1, ancho + 1)) + " |")
    lineas.append("|---" * (ancho + 1) + "|")
    for n, fila in enumerate(filas, 1):
        if any(fila[:ancho]):
            celdas = (fila + [""] * ancho)[:ancho]
            lineas.append(f"| {n} | " + " | ".join(celdas) + " |")
    if formulas:
        lineas.append("")
        lineas.append(f"Fórmulas de la hoja «{nombre}» (celda: fórmula → valor calculado):")
        for celda, formula, valor in formulas[:MAX_FORMULAS]:
            lineas.append(f"- {celda}: {formula} → {valor or '(sin valor calculado: abra y guarde el archivo en Excel)'}")
        if len(formulas) > MAX_FORMULAS:
            lineas.append(f"- … y {len(formulas) - MAX_FORMULAS} fórmulas más.")
    return "\n".join(lineas)


def _excel(ruta: Path) -> str:
    from openpyxl import load_workbook

    # Dos lecturas: con los valores calculados que Excel guarda y con las fórmulas.
    valores = load_workbook(ruta, read_only=True, data_only=True)
    formulas = load_workbook(ruta, read_only=True, data_only=False)
    hojas = []
    try:
        for hoja_v in valores.worksheets:
            hoja_f = formulas[hoja_v.title]
            # Se formatea cada valor con el formato de su celda (fechas, porcentajes…).
            filas = [[_valor(c.value, getattr(c, "number_format", "")) for c in fila]
                     for fila in hoja_v.iter_rows()]
            lista = []
            for n, fila in enumerate(hoja_f.iter_rows(values_only=True), 1):
                for c, contenido in enumerate(fila, 1):
                    if isinstance(contenido, str) and contenido.startswith("="):
                        calculado = filas[n - 1][c - 1] if n <= len(filas) and c <= len(filas[n - 1]) else ""
                        lista.append((f"{_letra(c)}{n}", contenido, calculado))
            oculta = " (hoja oculta)" if getattr(hoja_v, "sheet_state", "visible") != "visible" else ""
            hojas.append(_tabla(hoja_v.title, filas, lista, oculta))
    finally:
        valores.close()
        formulas.close()
    return f"Libro de Excel con {len(hojas)} hoja(s).\n\n" + "\n\n".join(hojas)


def _excel_antiguo(ruta: Path) -> str:
    import xlrd

    libro = xlrd.open_workbook(ruta)
    hojas = []
    for hoja in libro.sheets():
        filas = []
        for n in range(hoja.nrows):
            fila = []
            for c in range(hoja.ncols):
                celda = hoja.cell(n, c)
                if celda.ctype == xlrd.XL_CELL_DATE:
                    fila.append(xlrd.xldate_as_datetime(celda.value, libro.datemode))
                elif celda.ctype == xlrd.XL_CELL_BOOLEAN:
                    fila.append(bool(celda.value))
                elif celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    fila.append(None)
                else:
                    fila.append(celda.value)
            filas.append(fila)
        oculta = " (hoja oculta)" if hoja.visibility else ""
        hojas.append(_tabla(hoja.name, filas, nota=oculta))
    return (f"Libro de Excel (formato .xls antiguo) con {len(hojas)} hoja(s). Este formato no permite "
            "leer las fórmulas; solo los valores.\n\n" + "\n\n".join(hojas))


def _csv(ruta: Path) -> str:
    datos = ruta.read_bytes()
    # Excel en español suele exportar CSV en Windows-1252 y con «;» como separador.
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            texto = datos.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecto = csv.Sniffer().sniff(texto[:5000], delimiters=";,\t|")
    except csv.Error:
        dialecto = csv.excel
    filas = [list(f) for f in csv.reader(io.StringIO(texto), dialecto)]
    return _tabla(ruta.stem, filas)
